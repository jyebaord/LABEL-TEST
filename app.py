import io
import pandas as pd
import streamlit as st
from docx import Document
from docx.shared import Inches

# 設定網頁標題與寬版面
st.set_page_config(
    page_title="出貨外箱標籤與裝箱明細系統", page_icon="📦", layout="wide"
)

st.title("📦 出貨外箱標籤與裝箱明細網頁轉換系統")
st.markdown("請將 Excel 出貨明細檔案拖曳或上傳至下方，系統將自動為您排版。")

# 1. 檔案拖曳與上傳區
uploaded_file = st.file_uploader(
    "選擇或拖曳 Excel 檔案 (.xlsx, .xls)", type=["xlsx", "xls"]
)

if uploaded_file is not None:
  try:
    # 讀取 Excel 檔案 (略過前3行表頭)
    df = pd.read_excel(uploaded_file, skiprows=3)

    # 過濾有效資料
    df = df.dropna(subset=["品號"])
    df = df[df["客戶簡稱"].notna() & (df["客戶簡稱"] != "合計")]

    st.success(f"成功讀取檔案！共載入 {len(df)} 筆明細資料。")

    # 2. 網頁表格即時預覽
    st.subheader("📋 載入的明細資料預覽")
    st.dataframe(df, use_container_width=True)

    # 3. 匯出 Word 區塊
    st.subheader("📥 匯出 Word 標籤檔案")

    if st.button("一鍵生成 Word 檔案", type="primary"):
      doc = Document()

      # 設定頁面邊距
      for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

      # 上半部：裝箱明細總表
      doc.add_heading("出貨外箱標籤 & 包裝明細裝箱明細", level=1)
      table = doc.add_table(rows=1, cols=5)
      table.style = "Table Grid"
      hdr_cells = table.rows[0].cells
      headers = ["C/N", "出貨日期", "捷寶料號", "出貨數量", "銷貨單號"]
      for i, header_text in enumerate(headers):
        hdr_cells[i].text = header_text
        for paragraph in hdr_cells[i].paragraphs:
          for run in paragraph.runs:
            run.font.bold = True

      c_n = 1
      for _, row in df.iterrows():
        row_cells = table.add_row().cells
        date_str = (
            pd.to_datetime(row["出貨日期"]).strftime("%Y/%m/%d")
            if pd.notnull(row["出貨日期"])
            else ""
        )
        row_cells[0].text = str(c_n)
        row_cells[1].text = date_str
        row_cells[2].text = str(row["品號"])
        row_cells[3].text = str(row["出貨數量[訂]"])
        row_cells[4].text = str(row["出貨單號"])
        c_n += 1

      doc.add_paragraph("\n")

      # 下半部：各品項外箱標籤卡片
      doc.add_heading("外箱標籤預覽與列印", level=1)
      for _, row in df.iterrows():
        client = str(row["客戶簡稱"])
        part_no = str(row["品號"])
        qty = str(row["出貨數量[訂]"])
        doc_no = str(row["出貨單號"])

        qty_display = f"{qty} /pcs"
        if ".070CU" in part_no or ".047CU" in part_no:
          qty_display = f"{qty}/M =>每支1M共*{qty}EA(整直)"

        p = doc.add_paragraph()
        p.add_run("客戶簡稱: ").bold = True
        p.add_run(f"{client}\n")
        p.add_run("捷寶產品料號: ").bold = True
        p.add_run(f"{part_no}\n")
        p.add_run("產品出貨數量: ").bold = True
        p.add_run(f"{qty_display}\n")
        p.add_run("捷寶出貨單號: ").bold = True
        p.add_run(f"{doc_no}\n")
        p.add_paragraph("-" * 40)

      # 儲存至記憶體緩衝區
      buffer = io.BytesIO()
      doc.save(buffer)
      buffer.seek(0)

      # 提供網頁下載按鈕
      st.download_button(
          label="點擊下載 Word 檔案 (.docx)",
          data=buffer,
          file_name="出貨標籤與裝箱明細.docx",
          mime=(
              "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
          ),
      )
      st.success("Word 檔案已準備完成，請點擊上方按鈕下載！")

  except Exception as e:
    st.error(f"處理檔案時發生錯誤：{str(e)}")